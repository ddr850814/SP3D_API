# -*- coding: utf-8 -*-
"""管道材料·焊接量统计工具（matpy）：CLB.vb / weldpy 的 Python/ttkbootstrap 复刻。

SP3D 管道材料清单 SQL 直连：注册表/认证 ini 自动凭据（同 weldpy），
四类管道对象（直管/管件/特殊件/仪表）按 (管线, PartNo) 聚合，直管
数量=长度（米），CDB 目录补属性后出「管道材料」单表 Excel。

焊接量统计（自 weldpy 整体迁入）：规格表切割长度/材料组按项目记忆
（spec_mem.json，兼容旧 WeldWeld 文件），焊口链 RDB 直连 + 直管聚合。

界面与交互骨架（ttkbootstrap cosmo + 线程 queue/after 轮询）：报表下拉
切换布局——焊接量报表显示规格表与焊接选项，其余报表为材料报表布局。
"""
import datetime
import os
import queue
import sys
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox

import ttkbootstrap as ttk

import db
import dialogs
import mat_core
import snapshots
import weld_core
from openpyxl import Workbook, load_workbook

BASE = Path(__file__).resolve().parent
THEME = "cosmo"
POLL_MS = 80


def _build_stamp():
    """exe/源码文件的修改时间，作为构建标识显示在标题栏（exe 打包后即固定）。"""
    p = sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__)
    try:
        return datetime.datetime.fromtimestamp(os.path.getmtime(p)).strftime("%m-%d %H:%M")
    except OSError:
        return "?"


def _fmt_cut(v):
    """切割长度统一为紧凑数字文本（6.0 -> 6），非法值原样返回。"""
    try:
        return f"{float(str(v).strip()):g}"
    except (TypeError, ValueError):
        return str(v).strip()


# 「隐含件」级联子下拉的四种报表（XOwnsImpliedItems 四类目的地：
# 螺栓/垫片/管夹/焊口，焊口即焊接量统计）
_IMPLIED_KINDS = ("螺栓", "垫片", "管夹", "焊接量统计")


def _implied_writer(report):
    """隐含件报表名 → excel_out 写表函数（螺栓/垫片/管夹三档 by_line）。"""
    import excel_out
    return {"螺栓": excel_out.write_bolts,
            "垫片": excel_out.write_gaskets,
            "管夹": excel_out.write_clamps}[report]


def _geo_ok(val, default):
    """校验 "宽x高" 尺寸串，过小或非法回退默认值。"""
    try:
        w, h = str(val).lower().split("x")
        w, h = int(w), int(h)
        if w >= 400 and h >= 300:
            return f"{w}x{h}"
    except (ValueError, AttributeError):
        pass
    return default


class App(ttk.Window):
    def __init__(self):
        super().__init__(themename=THEME)
        self.title(f"SP3D 管道报表工具（独立版 构建 {_build_stamp()}）")
        # 窗口高清图标（dialogs.set_app_icon 内含 Tk 图标方案的实测结论）
        dialogs.set_app_icon(self)
        self.geometry("660x420")
        self.minsize(560, 360)

        self.cfg = db.load_config()
        self.plants = []
        self.plant = None
        self.sys_names = []
        self.sys_oids = []
        self.vol_oids = []   # 过滤器 Volume 条件引用的体积对象（材料表交集用）
        self.vol_boxes = []  # Planes 条件的平面组围合盒 [x0,y0,z0,x1,y1,z1]
        self.dwg_folder_names = []   # ISO 图纸目录勾选（独立于系统过滤器）
        self.dwg_folder_oids = []
        self.groups = []             # 项目材料组列表（焊接量报表，Schema CodelistValueView）
        self.all_specs = []          # 数据库拉取的完整规格列表（用于计算删除状态记忆）
        self._busy = 0            # 重叠加载计数：光标/进度条按计数配对，防止卡死
        self._tokens = {"plants": 0, "specs": 0, "stats": 0}
        self._last_rows = None    # 上次采集结果缓存（保存取消/失败后免重查）
        self._last_key = None     # 缓存对应的 (项目, 范围) 键
        self._last_detail = None  # 隐含件（螺栓/垫片/管夹）缓存 (报表名, rows)
        self._last_welds = None   # 焊接量统计缓存（保存取消/失败后免重算）
        self._last_opts = None
        self._last_plant = ""
        self._pending_opts = None
        self._specs_plant = ""    # 规格表当前对应的项目代码（防把旧项目数据记错）
        self._reread = False      # 重读数据库标记：本次规格加载忽略本地记忆
        self._q = queue.Queue()
        # 窗口尺寸：优先用户上次手动调整值；无记忆则按内容自适应
        # （旧版强制 660x420/660x660 的存档值视为无记忆，避免继续沿用不合适的高度）
        geo_mat = _geo_ok(self.cfg.get("geo_mat"), "")
        geo_pid = _geo_ok(self.cfg.get("geo_pid"), "")
        geo_weld = _geo_ok(self.cfg.get("geo_weld"), "")
        self._geo = {"mat": "" if geo_mat in ("", "660x420") else geo_mat,
                     "pid": "" if geo_pid in ("", "660x420") else geo_pid,
                     "weld": "" if geo_weld in ("", "660x660") else geo_weld}
        self._geo_applied = {}   # 各模式最近一次程序设置的尺寸（防启动期误记）

        self.style.configure(".", font=("Microsoft YaHei", 10))
        # 规格表行高加大（编辑器点击更准），仅作用于焊接量规格表，不影响其他树
        self.style.configure("Weld.Treeview", rowheight=30)
        self._build_ui()
        self.bind("<Configure>", self._on_configure)
        self._on_report_changed()  # 恢复上次报表并同步选项可用态
        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self.after(POLL_MS, self._poll)
        self._refresh_conn_state()
        self.after(0, self._auto_connect)

    # ---------------- UI ----------------
    def _build_ui(self):
        root = ttk.Frame(self, padding=(16, 14, 16, 14))
        root.pack(fill="both", expand=True)

        # 顶部：连接状态（凭据全部来自认证 ini，无手动设置入口）
        self.lbl_conn = tk.Label(root, anchor="w", font=("Microsoft YaHei", 10))
        self.lbl_conn.pack(fill="x", pady=(0, 8))

        # 报表选择（多报表：管道材料 / 隐含件（级联子下拉：螺栓/垫片/
        # 管夹/焊接量统计，四种隐含件共用一个入口）/ 仪表 / … ）
        row0 = ttk.Frame(root)
        row0.pack(fill="x")
        ttk.Label(row0, text="报表:", width=8, anchor="e").pack(side="left")
        self.cmb_report = ttk.Combobox(
            row0, state="readonly", width=14,
            values=["管道材料", "隐含件", "仪表", "管道属性",
                    "PID/应力号", "PID关联", "ISO图纸"])
        # 记忆恢复：隐含件四种按子下拉展开，其余直接设主下拉
        saved_rep = self.cfg.get("report", "管道材料")
        if saved_rep in _IMPLIED_KINDS:
            self.cmb_report.set("隐含件")
        else:
            self.cmb_report.set(saved_rep if saved_rep != "隐含件"
                                else "管道材料")
        self.cmb_report.bind("<<ComboboxSelected>>", self._on_report_changed)
        self.cmb_report.pack(side="left", padx=(4, 0), ipady=1)
        # 隐含件子下拉（主选「隐含件」时显示）
        self.cmb_implied = ttk.Combobox(
            row0, state="readonly", width=12, values=list(_IMPLIED_KINDS))
        self.cmb_implied.set(saved_rep if saved_rep in _IMPLIED_KINDS
                             else "螺栓")
        self.cmb_implied.bind("<<ComboboxSelected>>", self._on_report_changed)

        # 项目选择
        row1 = ttk.Frame(root)
        row1.pack(fill="x", pady=(8, 0))
        ttk.Label(row1, text="项目:", width=8, anchor="e").pack(side="left")
        self.cmb_plant = ttk.Combobox(row1, state="readonly")
        self.cmb_plant.bind("<<ComboboxSelected>>", self._on_plant_changed)
        self.btn_refresh = ttk.Button(row1, text="刷新", width=8,
                                      command=self._refresh_plants)
        self.btn_refresh.configure(bootstyle="secondary-outline")
        self.cmb_plant.pack(side="left", fill="x", expand=True,
                            padx=(4, 6), ipady=1)
        self.btn_refresh.pack(side="left", ipady=2)

        # 过滤器范围
        row2 = ttk.Frame(root)
        row2.pack(fill="x", pady=(8, 0))
        self.btn_sys = ttk.Button(row2, text="选择过滤器...", width=14,
                                  command=self._open_filter_dialog)
        self.btn_sys.configure(bootstyle="secondary-outline")
        self.btn_sys.pack(side="left", ipady=2)
        self.lbl_sys = tk.Label(row2, text="未选择（全部系统）", anchor="w",
                                fg="#6b7280", font=("Microsoft YaHei", 10))
        self.lbl_sys.pack(side="left", fill="x", expand=True, padx=(8, 0), pady=(6, 0))

        # 说明 + 选项（材料报表专属；焊接量报表下隐藏，见 _layout_report_ui）
        row3 = ttk.Frame(root)
        row3.pack(fill="x", pady=(12, 0))
        self.row3_mat = row3
        self.chk_by_line = ttk.BooleanVar(value=self.cfg.get("by_line", False))
        self.chk_by_line_cb = ttk.Checkbutton(
            row3, text="按管线号分组（不勾选则按商品代码汇总）",
            variable=self.chk_by_line, command=self._mem_opts)
        self.chk_by_line_cb.pack(side="left")
        ttk.Label(row3, text="材料范围:").pack(side="left", padx=(16, 4))
        self.cmb_hidden = ttk.Combobox(
            row3, state="readonly", width=10, values=["排除隐藏", "仅隐藏"])
        self.cmb_hidden.set("仅隐藏" if self.cfg.get("hidden_only") else "排除隐藏")
        self.cmb_hidden.bind("<<ComboboxSelected>>",
                             lambda e: self._on_hidden_changed())
        self.cmb_hidden.pack(side="left")
        note = tk.Label(root, anchor="w", justify="left", fg="#6b7280",
                        font=("Microsoft YaHei", 9),
                        text="统计范围：直管/管件/特殊件/仪表四类管道件；直管数量为长度（米），"
                             "其余为件数；特殊件与仪表按位号单列。\n"
                             "管道材料输出字段含：CC码、部件类型、管道等级/名称、操作与设计压力温度、"
                             "保温、单位重量、单元/区域名称（按管线号分组时）；"
                             "螺栓/垫片/管夹（隐含件子下拉）同样支持按管线号分组"
                             "（逐件明细）与材料范围选项，不勾选分组时按商品代码汇总数量；"
                             "管夹商品代码取自管道等级定义（按通径段匹配）。")
        note.pack(fill="x", pady=(8, 0))
        self.note = note

        # ---- 焊接量报表专属控件（weldpy 迁入；初始不 pack，
        #      由 _layout_report_ui 按报表显隐，避免 pack_forget 后重排错位）----
        # 规格表标题行（左标题右按钮，中间自然留白不会重叠）
        row_t = ttk.Frame(root)
        title = ttk.Label(row_t, text="规格切割长度与材料组",
                          font=("Microsoft YaHei", 10, "bold"))
        title.pack(side="left", pady=(6, 0))
        self.btn_export = ttk.Button(row_t, text="导出 xlsx", width=9,
                                     command=self._export_specs)
        self.btn_import = ttk.Button(row_t, text="导入 xlsx", width=9,
                                     command=self._import_specs)
        self.btn_reread = ttk.Button(row_t, text="重读数据库", width=10,
                                     command=self._reread_specs)
        for b in (self.btn_export, self.btn_import, self.btn_reread):
            b.configure(bootstyle="secondary-outline")
            b.pack(side="right", padx=(4, 0), ipady=2)
        self.row_t = row_t

        # 规格表（iid=规格名，唯一；材料组列按分组方式显示/隐藏）。
        # 注意 master 必须直接用容器 box：pack(in_=非父容器) 属边缘用法，
        # 部分环境下 Treeview 不重绘（整片空白连表头都没有）
        box = ttk.Frame(root)
        cols = ("spec", "cut", "grp")
        self.grid = ttk.Treeview(box, columns=cols, show="headings",
                                 selectmode="browse", style="Weld.Treeview")
        self.grid.heading("spec", text="规格")
        self.grid.heading("cut", text="切割长度")
        self.grid.heading("grp", text="材料组")
        # 规格/材料组列随窗口拉伸分享多余宽度，切割长度列固定
        self.grid.column("spec", width=240, minwidth=160, anchor="w", stretch=True)
        self.grid.column("cut", width=110, minwidth=110, anchor="center", stretch=False)
        self.grid.column("grp", width=150, minwidth=110, anchor="w", stretch=True)
        self.grid.configure(displaycolumns=("#all"))
        sb = ttk.Scrollbar(box, orient="vertical", command=self.grid.yview)
        self.grid.configure(yscrollcommand=sb.set)
        self.grid.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        self.grid.bind("<Double-1>", self._on_grid_double)
        self.grid.bind("<F2>", self._on_grid_edit_key)
        self.grid.bind("<Return>", self._on_grid_edit_key)
        self.grid.bind("<Delete>", self._on_grid_delete)
        self.grid.bind("<Button-3>", self._on_grid_popup)
        self.box = box

        # 焊接选项行
        row3_weld = ttk.Frame(root)
        ttk.Label(row3_weld, text="分组方式:", width=8, anchor="e").pack(side="left")
        self.cmb_met = ttk.Combobox(row3_weld, state="readonly", width=12,
                                    values=["不分组", "按材料等级", "按材料组"])
        try:
            self.cmb_met.current(int(self.cfg.get("met", 0)))
        except (ValueError, tk.TclError):
            self.cmb_met.current(0)
        self.cmb_met.bind("<<ComboboxSelected>>", lambda _e: self._on_met_changed())
        self.cmb_met.pack(side="left", padx=(4, 16))
        self.chk_thk = ttk.BooleanVar(value=self.cfg.get("thk", False))
        self.chk_skip_pipes = ttk.BooleanVar(value=self.cfg.get("skip_pipes", False))
        ttk.Checkbutton(row3_weld, text="忽略壁厚", variable=self.chk_thk,
                        command=self._mem_weld_opts).pack(side="left", padx=(0, 12))
        ttk.Checkbutton(row3_weld, text="忽略直管焊缝", variable=self.chk_skip_pipes,
                        command=self._mem_weld_opts).pack(side="left")
        self.row3_weld = row3_weld

        # PID 关联对象类别行（原版 frmReportCorrelate 过滤器提示的 7 类
        # 正选；初始不 pack，由 _layout_report_ui 按报表显隐）
        row3_pid = ttk.Frame(root)
        ttk.Label(row3_pid, text="对象类别:", width=8, anchor="e").pack(side="left")
        saved = self.cfg.get("pid_kinds")
        kinds_on = set(mat_core.PID_KIND_LABELS) if saved is None else set(saved)
        self.pid_kind_vars = {}
        for kind, label in mat_core.PID_KIND_LABELS.items():
            var = ttk.BooleanVar(value=kind in kinds_on)
            self.pid_kind_vars[kind] = var
            ttk.Checkbutton(row3_pid, text=label, variable=var,
                            command=self._mem_pid_kinds).pack(
                side="left", padx=(0, 4))
        self.row3_pid = row3_pid

        # 统计行
        row4 = ttk.Frame(root)
        row4.pack(fill="x", pady=(14, 0))
        self.row4 = row4
        self.btn_run = ttk.Button(row4, text="统计并生成 Excel",
                                  bootstyle="primary", command=self._run_stats)
        self.btn_run.configure(state="disabled")
        self.btn_run.pack(side="left", ipady=5, ipadx=8)
        self.btn_snap = ttk.Button(row4, text="存为快照",
                                   bootstyle="secondary-outline",
                                   command=self._save_snapshot)
        self.btn_snap.pack(side="left", ipady=5, ipadx=6, padx=(10, 0))
        self.btn_cmp = ttk.Button(row4, text="快照比较...",
                                  bootstyle="secondary-outline",
                                  command=self._open_snapshot_dialog)
        self.btn_cmp.pack(side="left", ipady=5, ipadx=6, padx=(8, 0))
        self.progress = ttk.Progressbar(row4, mode="indeterminate")
        self.progress.pack(side="left", fill="x", expand=True,
                           padx=(12, 0), pady=(13, 13))

    # ---------------- 连接 ----------------
    def _refresh_conn_state(self):
        if self.cfg.get("server") and self.cfg.get("username"):
            self.lbl_conn.configure(
                text=f"已配置：{self.cfg['server']}（账号 {self.cfg['username']}）",
                fg="#1a7f37")
            self._refresh_plants()
        else:
            self.lbl_conn.configure(text="未找到认证文件（ini），无法自动连接",
                                    fg="#b54708")
            self.btn_run.configure(state="disabled")

    def _auto_connect(self):
        """启动自动连接：凭据缺失时自动弹出连接设置对话框。"""
        if not (self.cfg.get("server") and self.cfg.get("username")):
            messagebox.showinfo(
                "自动连接失败",
                "未找到有效的认证文件（ini）或服务器信息。\n"
                "请在连接设置中指定认证 ini 文件（服务器/Site 库也可在此核对）。",
                parent=self)
            self._open_settings()

    def _open_settings(self):
        cfg = dialogs.open_conn_dialog(self, self.cfg)
        if cfg is not None:
            self.cfg = cfg
            self._refresh_conn_state()

    @property
    def conn_info(self) -> dict:
        """cfg + 解密后的明文密码（供工作线程使用）。"""
        info = dict(self.cfg)
        info["password"] = db.decrypt_password(self.cfg.get("password_enc", ""))
        return info

    # ---------------- 后台任务派发 ----------------
    def _poll(self):
        """主线程事件泵：把工作线程的结果队列派发到对应回调。"""
        if not self.winfo_exists():
            return
        try:
            while True:
                kind, token, *rest = self._q.get_nowait()
                getattr(self, "_on_" + kind)(token, *rest)
        except queue.Empty:
            pass
        self.after(POLL_MS, self._poll)

    def _start(self, fn):
        threading.Thread(target=fn, daemon=True).start()

    # ---------------- Plant ----------------
    def _refresh_plants(self):
        if not (self.cfg.get("server") and self.cfg.get("username")):
            return
        self._busy_on()
        token = self._tokens["plants"] = self._tokens["plants"] + 1
        info = self.conn_info
        q = self._q

        def work():
            try:
                plants = db.list_plants(info["server"], info["site_db"],
                                        info["username"], info["password"],
                                        info.get("provider", "MSSQL"))
                q.put(("plants", token, plants))
            except Exception as ex:
                q.put(("plants_failed", token, str(ex)))
        self._start(work)

    def _on_plants(self, token, plants):
        self._busy_off()
        if token != self._tokens["plants"]:
            return  # 旧线程的晚到结果：仅回收 busy 计数，不覆盖界面
        self.plants = plants
        # 刷新只管项目列表：尽量停在原项目；启动首次加载用上次使用的项目
        cur = self._plant_key() or (self.cfg.get("last_plant") or "")
        labels = [f"{p['display']}（{p['code']}）" for p in plants]
        self.cmb_plant.configure(values=labels)
        idx = 0
        for i, p in enumerate(plants):
            if p.get("code") == cur:
                idx = i
                break
        if labels:  # 空列表 current(0) 会抛 TclError
            self.cmb_plant.current(idx)
        if plants:
            # 首次加载（尚无当前项目）或原项目已从库中移除时才切换
            sel = plants[idx]
            if not self.plant or sel.get("code") != self._plant_key():
                self._select_plant(sel)
        else:
            self.btn_run.configure(state="disabled")

    def _on_plants_failed(self, token, msg):
        self._busy_off()
        if token != self._tokens["plants"]:
            return
        if messagebox.askyesno(
                "自动连接失败",
                f"{msg}\n\n是否打开连接设置？（检查认证 ini / 服务器）",
                parent=self):
            self._open_settings()

    def _on_plant_changed(self, _e=None):
        idx = self.cmb_plant.current()
        if 0 <= idx < len(self.plants):
            self._select_plant(self.plants[idx])

    def _plant_key(self) -> str:
        return self.plant.get("code", "") if self.plant else ""

    def _select_plant(self, plant):
        # 切换前把当前项目的规格配置/过滤器范围写入记忆（焊接量报表）
        self._save_spec_mem()
        self.plant = plant
        # 记住本次使用的项目，下次启动自动选中
        self.cfg["last_plant"] = self._plant_key()
        db.save_config(self.cfg)
        if self._report() == "焊接量统计":
            # 焊接量：过滤器范围与规格配置存 spec_mem（weldpy 语义）
            f = (db.load_spec_mem().get(self._plant_key()) or {}).get("filter") or {}
            self.sys_names = [f["name"]] if f.get("name") else []
            self.sys_oids = [str(o) for o in (f.get("oids") or [])]
            self.vol_oids = [str(o) for o in (f.get("vol_oids") or [])]
            self.vol_boxes = [list(map(float, b)) for b in (f.get("vol_boxes") or [])]
            self._set_sys_label()
            self.btn_run.configure(state="normal")
            self._last_rows = None
            self._load_specs()
            return
        # 恢复该项目记忆的过滤器范围
        f = (self.cfg.get("plant_filters") or {}).get(self._plant_key()) or {}
        self.sys_names = [f["name"]] if f.get("name") else []
        self.sys_oids = [str(o) for o in (f.get("oids") or [])]
        self.vol_oids = [str(o) for o in (f.get("vol_oids") or [])]
        self.vol_boxes = [list(map(float, b)) for b in (f.get("vol_boxes") or [])]
        # 恢复该项目记忆的图纸目录勾选
        d = (self.cfg.get("plant_dwg_filters") or {}).get(self._plant_key()) or {}
        self.dwg_folder_names = [str(x) for x in (d.get("names") or [])]
        self.dwg_folder_oids = [str(o) for o in (d.get("oids") or [])]
        self._set_sys_label()
        self.btn_run.configure(state="normal")
        self._last_rows = None  # 换项目清缓存
        self._last_detail = None
        self._last_welds = None
        self._specs_plant = ""  # 项目已换，进入焊接量报表时需重载规格

    def _set_sys_label(self):
        """过滤器范围标签：已选过滤器用正常色，未选灰色弱化。
        ISO 图纸报表切换为图纸目录勾选状态。"""
        if self._report() == "ISO图纸":
            self.btn_sys.configure(text="选择图纸目录...")
            if self.dwg_folder_names:
                shown = "、".join(self.dwg_folder_names[:3])
                more = "…" if len(self.dwg_folder_names) > 3 else ""
                self.lbl_sys.configure(
                    text=f"图纸目录 {shown}{more}"
                         f"（{len(self.dwg_folder_oids)} 个节点）",
                    fg="#212529")
            else:
                self.lbl_sys.configure(text="未选择（全部图纸）", fg="#6b7280")
            return
        self.btn_sys.configure(text="选择过滤器...")
        if self.sys_names or self.vol_oids or self.vol_boxes:
            parts = []
            if self.sys_oids:
                parts.append(f"{len(self.sys_oids)} 个系统节点")
            n_vol = len(self.vol_oids) + len(self.vol_boxes)
            if n_vol:
                parts.append(f"{n_vol} 个体积")
            self.lbl_sys.configure(
                text=f"过滤器 {self.sys_names[0] if self.sys_names else ''}"
                     f"（{' + '.join(parts)}）",
                fg="#212529")  # 空串是非法颜色名，须给有效色
        else:
            self.lbl_sys.configure(text="未选择（全部系统）", fg="#6b7280")

    def _mem_filters(self):
        """过滤器选择按项目写入配置（防杀进程丢失）。"""
        pf = self.cfg.get("plant_filters")
        if not isinstance(pf, dict):
            pf = {}
        pf[self._plant_key()] = {"name": self.sys_names[0] if self.sys_names else "",
                                 "oids": self.sys_oids,
                                 "vol_oids": self.vol_oids,
                                 "vol_boxes": self.vol_boxes}
        self.cfg["plant_filters"] = pf
        db.save_config(self.cfg)

    def _on_close(self):
        if self._busy:
            # 后台任务仍在查询：直接销毁会让工作线程回调报错
            messagebox.showinfo("请稍候", "后台任务进行中，完成后再关闭窗口。",
                                parent=self)
            return
        self._save_spec_mem()  # 焊接量规格表/过滤器按项目落盘（非焊接模式自动跳过）
        # 窗口尺寸（按报表模式）随配置持久化，下次启动恢复
        self.cfg["geo_mat"] = self._geo["mat"]
        self.cfg["geo_pid"] = self._geo["pid"]
        self.cfg["geo_weld"] = self._geo["weld"]
        db.save_config(self.cfg)
        self.destroy()

    # ---------------- 系统范围 ----------------
    def _open_filter_dialog(self):
        if not self.plant:
            return
        if self._report() == "焊接量统计":
            # 焊接量（weldpy 语义）：过滤器选择存 spec_mem 而非 config.plant_filters
            result = dialogs.open_filter_dialog(
                self, self.conn_info, self.plant,
                plant_name=self.plant["display"] or self.plant["code"],
                last_name=self.sys_names[0] if self.sys_names else "")
            if result is not None:
                name, oids, vol_oids, vol_boxes = result
                self.sys_names = [name] if name else []
                self.sys_oids = list(oids)
                self.vol_oids = list(vol_oids)
                self.vol_boxes = [list(map(float, b)) for b in vol_boxes]
                self._set_sys_label()
                self._save_spec_mem()  # 过滤器选择立即按项目落盘（防杀进程丢失）
            return
        if self._report() == "ISO图纸":
            # 图纸目录树（Deliverable 层级，模型过滤器选不到）：
            # 数据在 MDB CoreRelations，经 RDB 主连接跨库读
            result = dialogs.open_dwg_folder_dialog(
                self, self.conn_info, self.plant,
                plant_name=self.plant["display"] or self.plant["code"],
                last_names=self.dwg_folder_names)
            if result is not None:
                names, oids = result
                self.dwg_folder_names = list(names)
                self.dwg_folder_oids = list(oids)
                self._set_sys_label()
                self._mem_dwg_filters()
                self._last_rows = None
            return
        # 过滤器存于 MDB；Db 为惰连接（真正连库发生在线程首次查询时）
        result = dialogs.open_filter_dialog(
            self, self.conn_info, self.plant,
            plant_name=self.plant["display"] or self.plant["code"],
            last_name=self.sys_names[0] if self.sys_names else "")
        if result is not None:
            name, oids, vol_oids, vol_boxes = result
            self.sys_names = [name] if name else []
            self.sys_oids = list(oids)
            self.vol_oids = list(vol_oids)
            self.vol_boxes = [list(map(float, b)) for b in vol_boxes]
            self._set_sys_label()
            self._mem_filters()
            self._last_rows = None  # 范围变化清缓存

    def _mem_dwg_filters(self):
        """图纸目录勾选按项目写入配置。"""
        pf = self.cfg.get("plant_dwg_filters")
        if not isinstance(pf, dict):
            pf = {}
        pf[self._plant_key()] = {"names": self.dwg_folder_names,
                                 "oids": self.dwg_folder_oids}
        self.cfg["plant_dwg_filters"] = pf
        db.save_config(self.cfg)

    # ---------------- 统计 ----------------
    def _report(self) -> str:
        """当前报表名：主选「隐含件」时取子下拉，其余取主下拉。"""
        if self.cmb_report.get() == "隐含件":
            return self.cmb_implied.get() or "螺栓"
        return self.cmb_report.get() or "管道材料"

    def _on_report_changed(self, _e=None):
        """切报表（主/子下拉共用）：记住选择并切换布局；分组/材料范围
        选项对管道材料与螺栓/垫片/管夹报表有效；范围按钮在 ISO 图纸
        报表切换为图纸目录；子下拉随主选「隐含件」显隐。"""
        if self.cmb_report.get() == "隐含件":
            self.cmb_implied.pack(side="left", padx=(6, 0), ipady=1)
        else:
            self.cmb_implied.pack_forget()
        old = self.cfg.get("report")
        self.cfg["report"] = self._report()
        db.save_config(self.cfg)
        if old == "焊接量统计" and self._report() != "焊接量统计":
            # 离开焊接量模式：网格仍属于当前项目，先固化规格/过滤器记忆
            self._save_spec_mem()
        self._layout_report_ui()
        material = self._report() == "管道材料"
        opt_like = self._report() in ("管道材料", "螺栓", "垫片", "管夹")
        state = "normal" if opt_like else "disabled"
        self.chk_by_line_cb.configure(state=state)
        self.cmb_hidden.configure(state="readonly" if opt_like else "disabled")
        self._set_sys_label()
        self._last_rows = None  # 报表变化清缓存
        self._last_detail = None
        if self._report() == "焊接量统计":
            self.btn_run.configure(text="统计焊接量并生成 Excel")
            # 进入焊接量模式：项目已选且规格表不属于当前项目时才加载
            # （网格数据仍在时直接沿用，未保存的编辑不丢失）
            if self.plant and self._specs_plant != self._plant_key():
                self._load_specs()
        else:
            self.btn_run.configure(
                text="统计并生成 Excel" if material
                else f"统计{self._report()}并生成 Excel")

    def _mode_key(self) -> str:
        """窗口几何分档：焊接量 / PID 关联 / 其余材料类报表各记一档，
        切报表时按各自内容自适应或恢复各自记忆高度。"""
        r = self._report()
        if r == "焊接量统计":
            return "weld"
        if r == "PID关联":
            return "pid"
        return "mat"

    def _on_configure(self, e):
        """记忆用户调整的窗口尺寸（按报表模式分开记；最大化与启动期不记）。"""
        if e.widget is not self or self.state() == "zoomed":
            return
        key = self._mode_key()
        if key not in self._geo_applied:
            return  # 首次布局前的事件不记，避免启动默认值污染记忆
        self._geo[key] = _geo_ok(f"{e.width}x{e.height}", self._geo[key])

    def _default_geo(self, key):
        """内容自适应默认高度：各模式紧凑贴合自身内容（底部不留空）；
        焊接量报表贴合内容 + 约 15 行规格表；均不超出屏幕。"""
        self.grid.configure(height=1)   # 排除规格表默认请求的 10 行高度
        self.update_idletasks()
        h = self.winfo_reqheight()
        if key == "weld":
            h += 15 * 30                # 预留 15 行规格表（行高 30）
            h = max(h, 520)
        else:
            h = max(h, 300)
        h = min(h, self.winfo_screenheight() - 120)
        self.grid.configure(height=10)
        return f"660x{h}"

    def _mode_geo(self, key):
        """取该模式应用的窗口尺寸：用户记忆优先，无记忆用自适应默认。"""
        if not self._geo[key]:
            self._geo[key] = self._default_geo(key)
        self._geo_applied[key] = self._geo[key]
        return self._geo[key]

    def _layout_report_ui(self):
        """按报表切换主窗口布局与几何：
        焊接量 → 规格表 + 焊接选项行；PID 关联 → 对象类别行；
        其余 → 材料选项行 + 说明。
        PID/焊接专属控件 pack 用 before=self.row4 锚点插入（row4 恒在），
        避免 pack_forget 后重新 pack 被追加到统计行之后。
        窗口高度按模式分档（weld/pid/mat）：有记忆恢复记忆值，
        无记忆按各自内容自适应（材料不留空、焊接给足规格表行数）。"""
        weld = self._report() == "焊接量统计"
        pid = self._report() == "PID关联"
        if weld:
            self.row3_mat.pack_forget()
            self.note.pack_forget()
            self.row3_pid.pack_forget()
            self.row_t.pack(fill="x", pady=(12, 4), before=self.row4)
            self.box.pack(fill="both", expand=True, pady=(4, 8), before=self.row4)
            self.row3_weld.pack(fill="x", pady=(4, 0), before=self.row4)
            self.minsize(560, 520)
            self.geometry(self._mode_geo("weld"))
        elif pid:
            self.row_t.pack_forget()
            self.box.pack_forget()
            self.row3_weld.pack_forget()
            self.row3_mat.pack_forget()
            self.note.pack_forget()
            self.row3_pid.pack(fill="x", pady=(12, 0), before=self.row4)
            self.minsize(560, 300)
            self.geometry(self._mode_geo("pid"))
        else:
            self.row_t.pack_forget()
            self.box.pack_forget()
            self.row3_weld.pack_forget()
            self.row3_pid.pack_forget()
            self.row3_mat.pack(fill="x", pady=(12, 0), before=self.row4)
            self.note.pack(fill="x", pady=(8, 0), before=self.row4)
            self.minsize(560, 300)
            self.geometry(self._mode_geo("mat"))

    def _hidden_only(self) -> bool:
        return self.cmb_hidden.get() == "仅隐藏"

    def _on_hidden_changed(self):
        self._mem_opts()
        self._last_rows = None  # 材料范围变化清缓存

    def _mem_opts(self):
        self.cfg["by_line"] = self.chk_by_line.get()
        self.cfg["hidden_only"] = self._hidden_only()
        db.save_config(self.cfg)

    def _mem_pid_kinds(self):
        """PID 关联对象类别勾选写入配置（防杀进程丢失）。"""
        self.cfg["pid_kinds"] = [k for k, v in self.pid_kind_vars.items()
                                 if v.get()]
        db.save_config(self.cfg)

    def _pid_kinds(self):
        """勾选的 PID 关联对象类别；全不勾视为不过滤（None）。"""
        if self._report() != "PID关联":
            return None
        chosen = frozenset(k for k, v in self.pid_kind_vars.items() if v.get())
        return chosen or None

    def _stats_key(self):
        return (self._report(), self._plant_key(), tuple(self.sys_oids),
                tuple(self.vol_oids), tuple(map(tuple, self.vol_boxes)),
                self._hidden_only())

    def _run_stats(self, snapshot: bool = False):
        if not self.plant:
            return
        if self._report() == "焊接量统计":
            self._run_weld_stats(snapshot)
            return
        rep = self._report()
        if snapshot and rep not in snapshots.SPECS:
            return
        # 选项与上次完全一致且上次结果未成功导出 → 直接复用缓存
        # （材料报表 + 隐含件三种明细，取消保存/导出失败后免重查）
        if (self._last_key == self._stats_key() and (
                (rep == "管道材料" and self._last_rows is not None)
                or (rep in ("螺栓", "垫片", "管夹")
                    and self._last_detail is not None
                    and self._last_detail[0] == rep))):
            if rep == "管道材料":
                if snapshot:
                    self._write_snapshot("管道材料", *self._last_rows)
                else:
                    self._on_stats_done(None, *self._last_rows)
            elif rep in ("螺栓", "垫片", "管夹"):
                if snapshot:
                    self._write_snapshot(rep, self._last_detail[1])
                else:
                    self._save_implied_flow(rep, self._last_detail[1])
            return
        self._busy_on()
        token = self._tokens["stats"] = self._tokens["stats"] + 1
        info, plant = self.conn_info, self.plant
        report = rep
        opts = mat_core.MatOptions(sys_oids=list(self.sys_oids),
                                   sys_names=list(self.sys_names),
                                   hidden_only=self._hidden_only(),
                                   dwg_folder_oids=list(self.dwg_folder_oids),
                                   pid_kinds=self._pid_kinds(),
                                   vol_oids=list(self.vol_oids),
                                   vol_boxes=[list(b) for b in self.vol_boxes])
        q = self._q
        snap = snapshot and report in snapshots.SPECS
        scope = self._scope_text() if snap else None
        by_line = self.chk_by_line.get()  # GUI 线程先抓（口径固化进快照）

        def work():
            try:
                db_r = db.open_db(info.get("provider", "MSSQL"), info["server"],
                                  plant["reports"], info["username"], info["password"])
                try:
                    if report == "螺栓":
                        rows = mat_core.extract_bolts(db_r, plant, opts)
                        if snap and rows:
                            q.put(("snap_done", token, snapshots.save(
                                "螺栓", plant["code"], scope, rows,
                                by_line=by_line)))
                        elif snap:
                            q.put(("stats_empty", token))
                        else:
                            q.put(("bolts_done", token, rows))
                    elif report == "垫片":
                        rows = mat_core.extract_gaskets(db_r, plant, opts)
                        if snap and rows:
                            q.put(("snap_done", token, snapshots.save(
                                "垫片", plant["code"], scope, rows,
                                by_line=by_line)))
                        elif snap:
                            q.put(("stats_empty", token))
                        else:
                            q.put(("gaskets_done", token, rows))
                    elif report == "管夹":
                        rows = mat_core.extract_clamps(db_r, plant, opts)
                        if snap and rows:
                            q.put(("snap_done", token, snapshots.save(
                                "管夹", plant["code"], scope, rows,
                                by_line=by_line)))
                        elif snap:
                            q.put(("stats_empty", token))
                        else:
                            q.put(("clamps_done", token, rows))
                    elif report == "仪表":
                        rows = mat_core.extract_instruments(db_r, plant, opts)
                        q.put(("instr_done", token, rows))
                    elif report == "管道属性":
                        rows = mat_core.extract_piping_runs(db_r, plant, opts)
                        q.put(("runs_done", token, rows))
                    elif report == "PID/应力号":
                        rows = mat_core.extract_pid_numbers(db_r, plant, opts)
                        q.put(("pid_done", token, rows))
                    elif report == "PID关联":
                        rows = mat_core.extract_pid_correlations(
                            db_r, plant, opts)
                        q.put(("corr_done", token, rows))
                    elif report == "ISO图纸":
                        rows = mat_core.extract_iso_drawings(db_r, plant, opts)
                        q.put(("dwg_done", token, rows))
                    else:
                        rows = mat_core.extract_materials(db_r, plant, opts)
                        if rows:
                            attrs = mat_core.load_catalog_attrs(
                                db_r, plant["catalog"], plant.get("schema", ""),
                                {r.partno for r in rows})
                            if snap:
                                q.put(("snap_done", token, snapshots.save(
                                    "管道材料", plant["code"], scope,
                                    rows, attrs, by_line=by_line)))
                            else:
                                q.put(("stats_done", token, rows, attrs))
                        else:
                            q.put(("stats_empty", token))
                finally:
                    db_r.close()
            except Exception as ex:
                q.put(("stats_failed", token, str(ex)))
        self._start(work)

    # ---------------- 快照（材料表比较，对齐 CLB 材料表比较功能）----
    def _scope_text(self) -> str:
        """快照元数据里的范围描述（识别用）。"""
        rng = "；".join(self.sys_names) if self.sys_names else "全部系统"
        n_vol = len(self.vol_oids) + len(self.vol_boxes)
        if n_vol:
            rng += f" ∩ {n_vol} 个体积"
        if self._report() == "焊接量统计":
            parts = ["范围：" + rng,
                     ("不分组", "按材料等级", "按材料组")[self.cmb_met.current()]]
            if self.chk_thk.get():
                parts.append("忽略壁厚")
            if self.chk_skip_pipes.get():
                parts.append("忽略直管焊缝")
            return "；".join(parts)
        return (f"范围：{rng}；材料范围："
                f"{'仅隐藏' if self._hidden_only() else '排除隐藏'}")

    def _save_snapshot(self):
        rep = self._report()
        key = "焊接量" if rep == "焊接量统计" else rep
        if key not in snapshots.SPECS:
            messagebox.showinfo(
                "提示", "该报表暂不支持快照。\n支持：管道材料 / 螺栓 / 垫片 / 管夹 / 焊接量统计。",
                parent=self)
            return
        if not self.plant:
            return
        if key == "焊接量":
            self._run_weld_stats(snapshot=True)
        else:
            self._run_stats(snapshot=True)

    def _write_snapshot(self, report, rows, attrs=None):
        """主线程写快照（统计缓存命中路径）。口径按当前界面勾选固化
        （缓存复用条件是选项一致，界面值即缓存数据的口径）。"""
        try:
            if report == "焊接量":
                by_line, weld_opts = None, {
                    "met": self.cmb_met.current(),
                    "thk": self.chk_thk.get(),
                    "skip_pipes": self.chk_skip_pipes.get()}
            else:
                by_line, weld_opts = self.chk_by_line.get(), None
            meta = snapshots.save(report, self.plant["code"],
                                  self._scope_text(), rows, attrs,
                                  by_line=by_line, weld_opts=weld_opts)
        except Exception as ex:
            messagebox.showerror("快照保存失败", str(ex), parent=self)
            return
        self._snap_saved_msg(meta)

    def _on_snap_done(self, token, meta):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        self._snap_saved_msg(meta)

    def _snap_saved_msg(self, meta):
        messagebox.showinfo(
            "快照已保存",
            f"快照已保存：\n{meta['path']}\n\n共 {len(meta['rows'])} 行。\n"
            "在「快照比较...」中选择两份同类型快照可比较差异。", parent=self)

    def _open_snapshot_dialog(self):
        rep = self._report()
        key = "焊接量" if rep == "焊接量统计" else rep
        # 只列出同类型、同项目、同口径（GUI 当前勾选）的快照；
        # 范围不限（扩大范围比增量是合理用法）
        flt = None
        if key in snapshots.SPECS and self.plant:
            flt = {"report": key, "plant": self.plant["code"]}
            if key == "焊接量":
                flt["weld_opts"] = {
                    "met": self.cmb_met.current(),
                    "thk": self.chk_thk.get(),
                    "skip_pipes": self.chk_skip_pipes.get()}
            else:
                flt["by_line"] = self.chk_by_line.get()
        dialogs.open_snapshot_dialog(self, self.cfg, flt)

    def _save_detail(self, rows, suffix: str, writer) -> bool:
        """明细通用保存流。返回 True=导出成功 / False=取消或失败。"""
        if not rows:
            messagebox.showwarning("提示", "采集结果为空。请检查过滤器范围。", parent=self)
            return False
        default_name = f"{self.plant['code']}_{suffix}.xlsx"
        start_dir = self.cfg.get("last_dir") or str(Path.home())
        path = filedialog.asksaveasfilename(
            parent=self, title="保存 Excel", initialdir=start_dir,
            initialfile=default_name, defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")])
        if not path:
            messagebox.showinfo(
                "提示",
                "已取消保存。采集结果已保留，再次点击「统计并生成 Excel」可直接导出"
                "（无需重新统计）。", parent=self)
            return False
        try:
            n = writer(rows, path)
            self.cfg["last_dir"] = str(Path(path).parent)
            db.save_config(self.cfg)
            messagebox.showinfo(
                "成功", f"文件保存成功！\n{path}\n\n"
                f"共 {n or len(rows)} 行{'汇总' if n is not None and n != len(rows) else '明细'}。",
                parent=self)
            return True
        except Exception as ex:
            messagebox.showerror(
                "导出失败",
                f"{ex}\n\n采集结果已保留，可再次点击「统计并生成 Excel」重新导出"
                "（无需重新统计）。", parent=self)
            return False

    def _save_implied_flow(self, report: str, rows) -> None:
        """隐含件（螺栓/垫片/管夹）保存流：带缓存复用（对齐材料表——
        选项一致时取消保存/导出失败后，再点统计直接重导出免重查）。"""
        self._last_detail = (report, rows)
        self._last_key = self._stats_key()
        by_line = self.chk_by_line.get()
        if self._save_detail(
                rows, report,
                lambda r, p: _implied_writer(report)(r, p, by_line)):
            self._last_detail = None  # 导出成功，清除缓存

    def _on_bolts_done(self, token, rows):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        self._save_implied_flow("螺栓", rows)

    def _on_gaskets_done(self, token, rows):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        self._save_implied_flow("垫片", rows)

    def _on_clamps_done(self, token, rows):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        self._save_implied_flow("管夹", rows)

    def _on_instr_done(self, token, rows):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        import excel_out
        self._save_detail(rows, "仪表", excel_out.write_instruments)

    def _on_runs_done(self, token, rows):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        import excel_out
        self._save_detail(rows, "管道属性", excel_out.write_piping_runs)

    def _on_pid_done(self, token, rows):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        import excel_out
        self._save_detail(rows, "PID应力号", excel_out.write_pid_numbers)

    def _on_corr_done(self, token, rows):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        import excel_out
        self._save_detail(rows, "PID关联", excel_out.write_pid_correlations)

    def _on_dwg_done(self, token, rows):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        import excel_out
        self._save_detail(rows, "ISO图纸", excel_out.write_iso_drawings)

    def _on_stats_done(self, token, rows, attrs=None):
        self._busy_off()
        if token is not None and token != self._tokens["stats"]:
            return  # 旧线程的晚到结果：仅回收 busy 计数，不覆盖界面
        if not rows:
            messagebox.showwarning("提示", "采集结果为空。请检查过滤器范围。", parent=self)
            return
        self._last_rows = (rows, attrs)
        self._last_key = self._stats_key()
        default_name = f"{self.plant['code']}_管道材料.xlsx"
        start_dir = self.cfg.get("last_dir") or str(Path.home())
        path = filedialog.asksaveasfilename(
            parent=self, title="保存 Excel", initialdir=start_dir,
            initialfile=default_name, defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")])
        if not path:
            messagebox.showinfo(
                "提示",
                "已取消保存。采集结果已保留，再次点击「统计并生成 Excel」可直接导出"
                "（无需重新统计）。", parent=self)
            return
        try:
            import excel_out
            excel_out.write_excel(rows, attrs or {}, self.chk_by_line.get(), path)
            self.cfg["last_dir"] = str(Path(path).parent)
            db.save_config(self.cfg)
            self._last_rows = None  # 导出成功，清除缓存
            messagebox.showinfo(
                "成功", f"文件保存成功！\n{path}\n\n共 {len(rows)} 行材料数据。",
                parent=self)
        except Exception as ex:
            messagebox.showerror(
                "导出失败",
                f"{ex}\n\n采集结果已保留，可再次点击「统计并生成 Excel」重新导出"
                "（无需重新统计）。", parent=self)

    def _on_stats_failed(self, token, msg):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        messagebox.showerror("统计失败", msg, parent=self)

    def _on_stats_empty(self, token):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        messagebox.showwarning("提示", "采集结果为空，且未设置过滤器范围。", parent=self)

    # ---------------- 规格表与焊接量统计（自 weldpy 迁入） ----------------
    def _mem_weld_opts(self):
        """焊接量勾选项（忽略壁厚/忽略直管焊缝）状态变化即时写入配置记忆。"""
        self.cfg["thk"] = self.chk_thk.get()
        self.cfg["skip_pipes"] = self.chk_skip_pipes.get()
        db.save_config(self.cfg)

    def _on_met_changed(self):
        show_grp = self.cmb_met.current() == 2
        self.grid.configure(displaycolumns=("#all") if show_grp else ("spec", "cut"))
        met = self.cmb_met.current()
        if self.cfg.get("met", 0) != met:  # weldpy 读而不存的遗留点：补持久化
            self.cfg["met"] = met
            db.save_config(self.cfg)

    def _load_specs(self):
        self._busy_on()
        token = self._tokens["specs"] = self._tokens["specs"] + 1
        info, plant = self.conn_info, self.plant
        q = self._q

        def work():
            db_c = None
            try:
                db_c = db.open_db(info.get("provider", "MSSQL"), info["server"],
                                  plant["catalog"], info["username"], info["password"])
                pairs, groups = weld_core.load_specs(db_c, plant["schema"])
                q.put(("specs", token, pairs, groups))
            except Exception as ex:
                q.put(("specs_failed", token, str(ex)))
            finally:
                if db_c is not None:
                    db_c.close()
        self._start(work)

    def _on_specs(self, token, pairs, groups):
        self._busy_off()
        if token != self._tokens["specs"]:
            return
        self.groups = groups
        self.all_specs = [spec for spec, _ in pairs]
        self._default_grp = {spec: grp for spec, grp in pairs}  # 导入新增行的材料组默认值
        self._specs_plant = self._plant_key()  # 网格数据归属落定，记忆才可写
        reread, self._reread = self._reread, False
        if reread:
            # 重读数据库：不套用本地记录（切割长度默认 6、材料组取库值、不滤已删）
            rows, deleted = {}, set()
        else:
            # 恢复该项目记忆的切割长度/材料组，跳过记忆中已删除的规格
            mem = db.load_spec_mem().get(self._plant_key()) or {}
            rows, deleted = mem.get("rows", {}), set(mem.get("deleted", []))
        for i in self.grid.get_children():
            self.grid.delete(i)
        for spec, grp in pairs:
            if spec in deleted:
                continue
            rec = rows.get(spec) or {}
            try:
                self.grid.insert("", "end", iid=spec, values=(
                    spec, _fmt_cut(rec.get("cut", "6")), rec.get("grp") or grp))
            except tk.TclError:
                continue  # 数据库出现重复规格名时跳过重复项（iid 冲突）
        self._on_met_changed()

    def _on_specs_failed(self, token, msg):
        self._busy_off()
        if token != self._tokens["specs"]:
            return
        self._reread = False  # 加载失败不残留重读标记，避免下次刷新误跳过记忆
        messagebox.showerror("加载项目失败", msg, parent=self)

    def _on_grid_double(self, event):
        if self.grid.identify("region", event.x, event.y) != "cell":
            return
        col = self.grid.identify_column(event.x)
        row = self.grid.identify_row(event.y)
        if col == "#1":
            return  # 规格列只读
        if col == "#3" and self.cmb_met.current() != 2:
            return  # 材料组列隐藏时不可达，防御
        self._begin_edit(row, col)

    def _on_grid_edit_key(self, _e=None):
        row = self.grid.focus()
        if not row:
            return
        # 键盘编辑默认进切割长度列（材料组列用双击进）
        self._begin_edit(row, "#2")

    def _begin_edit(self, iid, col):
        bbox = self.grid.bbox(iid, col)
        if not bbox:  # 单元格滚动出视图
            return
        x, y, w, h = bbox
        cur = self.grid.set(iid, col)
        if col == "#3" and self.groups:
            ed = ttk.Combobox(self.grid, values=self.groups, state="readonly")
            ed.set(cur if cur in self.groups else self.groups[0])
        else:
            ed = ttk.Entry(self.grid)
            ed.insert(0, cur)
        ed.place(x=x, y=y, width=max(w, 90), height=h)
        ed.focus_set()
        if not isinstance(ed, ttk.Combobox):  # Combobox 继承自 Entry，须排除
            ed.select_range(0, "end")
        done = [False]

        def commit(_e=None):
            if done[0]:
                return
            done[0] = True
            val = ed.get().strip()
            if val:
                self.grid.set(iid, col, _fmt_cut(val) if col == "#2" else val)
            ed.destroy()

        def cancel(_e=None):
            if done[0]:
                return
            done[0] = True
            ed.destroy()

        ed.bind("<Return>", commit)
        ed.bind("<Escape>", cancel)
        ed.bind("<FocusOut>", commit)
        # Combobox 下拉展开时列表点击会先触发 FocusOut，延迟提交等选择生效
        if isinstance(ed, ttk.Combobox):
            ed.bind("<FocusOut>", lambda e: self.after(120, lambda: commit()))

    def _on_grid_delete(self, _e=None):
        for iid in self.grid.selection():
            self.grid.delete(iid)

    def _on_grid_popup(self, event):
        iid = self.grid.identify_row(event.y)
        if not iid:
            return
        self.grid.selection_set(iid)
        self.grid.focus(iid)
        m = tk.Menu(self, tearoff=0)
        m.add_command(label="删除选中规格行", command=self._on_grid_delete)
        try:
            m.tk_popup(event.x_root, event.y_root)
        finally:
            m.grab_release()

    def _save_spec_mem(self):
        """把当前规格表（切割长度/材料组/删除状态）与过滤器范围按项目写入本地记忆。

        仅当网格数据确实属于当前项目（_specs_plant 一致）时才写：
        材料报表下切项目/关窗口等调用点不会把旧项目的网格错记到新项目。
        """
        key = self._plant_key()
        if not key or self._busy or self._specs_plant != key:
            return
        rows = {}
        for iid in self.grid.get_children():
            spec, cut, grp = self.grid.item(iid, "values")
            rows[spec] = {"cut": _fmt_cut(cut) if cut else "6",
                          "grp": grp or ""}
        mem = db.load_spec_mem()
        mem[key] = {
            "rows": rows,
            "deleted": [s for s in self.all_specs if s not in rows],
            # 过滤器范围一并按项目记忆（切回项目免重选）
            "filter": {
                "name": self.sys_names[0] if self.sys_names else "",
                "oids": [str(o) for o in self.sys_oids],
                "vol_oids": [str(o) for o in self.vol_oids],
                "vol_boxes": [list(b) for b in self.vol_boxes],
            },
        }
        db.save_spec_mem(mem)

    def _export_specs(self):
        if not len(self.grid.get_children()):
            messagebox.showinfo("提示", "规格表为空，无可导出内容。", parent=self)
            return
        default = f"{self._plant_key() or '规格表'}_切割长度.xlsx"
        last_dir = self.cfg.get("last_dir") or "."
        path = filedialog.asksaveasfilename(
            parent=self, title="导出规格表",
            initialdir=last_dir, initialfile=default,
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "规格"
        ws.append(["规格", "切割长度", "材料组"])
        count = 0
        for iid in self.grid.get_children():
            spec, cut, grp = self.grid.item(iid, "values")
            try:
                cut_v = float(cut) if cut and str(cut).strip() else ""
            except ValueError:
                cut_v = cut
            ws.append([spec, cut_v, grp or ""])
            count += 1
        try:
            wb.save(path)
        except OSError as e:
            messagebox.showwarning("导出失败", f"无法写入文件：\n{e}", parent=self)
            return
        messagebox.showinfo("完成", f"已导出 {count} 行到：\n{path}", parent=self)

    def _import_specs(self):
        """导入 xlsx：使规格表与文件完全一致（新增/更新/删除）。

        - 文件中的规格：表格没有则新增，有则覆盖（切割长度空=默认 6，
          材料组不在列表/为空=取数据库默认值）
        - 表格中有而文件中没有的规格：删除（可「重读数据库」还原）
        - 文件中不属于当前项目的规格：跳过
        """
        path = filedialog.askopenfilename(
            parent=self, title="导入规格表",
            initialdir=self.cfg.get("last_dir", "."),
            filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        try:
            ws = load_workbook(path, read_only=True, data_only=True).active
            data = {}
            dups, bad_cuts, bad_grps = [], [], []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not str(row[0] or "").strip():
                    continue
                spec = str(row[0]).strip()
                cut = row[1] if len(row) > 1 and row[1] is not None else ""
                grp = row[2] if len(row) > 2 and row[2] is not None else ""
                grp = str(grp).strip()
                if spec in data:
                    dups.append(spec)
                # 切割长度校验（与统计口径一致）：必须为正数；
                # 空值允许（导入时取默认 6），非法值重置为默认并记录
                try:
                    ok = float(str(cut).strip()) > 0
                except (TypeError, ValueError):
                    ok = not str(cut).strip()
                if not ok:
                    bad_cuts.append((spec, cut))
                    cut = ""
                # 材料组校验：必须是项目材料组列表（Schema CodelistValueView）
                # 里的值；不在列表将回退该规格的数据库默认值
                if grp and self.groups and grp not in self.groups:
                    bad_grps.append((spec, grp))
                data[spec] = (_fmt_cut(cut), grp)
        except Exception as e:
            messagebox.showwarning("导入失败", f"无法读取文件：\n{e}", parent=self)
            return
        warn = ""
        if bad_cuts:
            shown = "、".join(f"{s}({v})" for s, v in bad_cuts[:5])
            more = " 等" if len(bad_cuts) > 5 else ""
            warn += f"\n· {len(bad_cuts)} 个无效切割长度将设为默认 6：{shown}{more}"
        if bad_grps:
            shown = "、".join(f"{s}[{g}]" for s, g in bad_grps[:5])
            more = " 等" if len(bad_grps) > 5 else ""
            warn += (f"\n· {len(bad_grps)} 个材料组不在项目列表，将用数据库默认值："
                     f"{shown}{more}")
        if dups:
            warn += f"\n· {len(dups)} 个重复规格按最后一行为准"
        if not messagebox.askyesno(
                "导入规格表",
                "导入后规格表将与文件内容完全一致：\n"
                "· 文件中的规格：新增或更新\n"
                "· 表格中有而文件中没有的规格：删除\n\n"
                "结果会立即保存到本地记录（删错了可点「重读数据库」还原）。"
                + warn + "\n\n是否继续？", parent=self):
            return
        defaults = getattr(self, "_default_grp", {})
        existing = set(self.grid.get_children())
        added = updated = deleted = skipped = 0
        for spec, (cut, grp) in data.items():
            if spec not in self.all_specs:
                skipped += 1
                continue
            cut_v = cut if cut else "6"
            grp_v = grp if grp in self.groups else defaults.get(spec, "")
            if spec in existing:
                self.grid.set(spec, "cut", cut_v)
                self.grid.set(spec, "grp", grp_v)
                updated += 1
            else:
                try:
                    self.grid.insert("", "end", iid=spec,
                                     values=(spec, cut_v, grp_v))
                    added += 1
                except tk.TclError:
                    pass  # 防御：极端情况下 iid 冲突
        for iid in existing:
            if iid not in data:
                self.grid.delete(iid)
                deleted += 1
        self._save_spec_mem()
        msg = f"更新 {updated}，新增 {added}，删除 {deleted}。"
        if skipped:
            msg += f"\n文件中 {skipped} 个规格当前项目不存在，已跳过。"
        messagebox.showinfo("完成", msg, parent=self)

    def _reread_specs(self):
        """忽略本地记忆，从数据库重读规格表（切割长度/材料组回到库值）。"""
        if not self.plant or self._busy:
            return  # 后台任务进行中不允许叠加加载
        if not messagebox.askyesno(
                "重读数据库",
                "将从数据库重新读取本项目规格清单：\n"
                "切割长度恢复默认 6，材料组取数据库值，已删规格还原。\n\n"
                "之后统计/切项目/关闭窗口会按重读后的表格更新本地记录。是否继续？",
                parent=self):
            return
        self._reread = True
        self._load_specs()

    def _run_weld_stats(self, snapshot: bool = False):
        clr, mg, bad = {}, {}, []
        for iid in self.grid.get_children():
            spec, cut, grp = self.grid.item(iid, "values")
            if cut is not None and str(cut).strip():
                try:
                    v = float(cut)
                except ValueError:
                    bad.append((spec, cut, "不是数字"))
                else:
                    # 非正数会导致管子聚合 SQL 除零/负除，必须拒绝
                    if v > 0:
                        clr[spec] = v
                    else:
                        bad.append((spec, cut, "必须为正数"))
            if grp:
                mg[spec] = grp
        if bad:
            shown = "\n".join(f"规格 {s}：切割长度 {v}（{why}）" for s, v, why in bad[:5])
            more = f"\n...等共 {len(bad)} 项" if len(bad) > 5 else ""
            messagebox.showwarning(
                "切割长度无效",
                "以下规格未纳入统计，请修正后重新统计：\n" + shown + more,
                parent=self)
        if not clr:
            messagebox.showwarning("提示", "请至少配置一个有效的切割长度。", parent=self)
            return
        self._save_spec_mem()  # 统计时同步固化规格记忆
        opts = weld_core.StatsOptions(
            clr=clr, mg=mg, met=self.cmb_met.current(),
            thk=self.chk_thk.get(),
            sys_names=self.sys_names, sys_oids=self.sys_oids,
            vol_oids=list(self.vol_oids),
            vol_boxes=[list(b) for b in self.vol_boxes],
            skip_pipes=self.chk_skip_pipes.get())
        self._pending_opts = opts
        # 选项与上次完全一致且上次结果未成功导出 → 直接复用缓存
        # （保存对话框点取消/写文件失败后，再点统计不必重算几十秒）
        if (self._last_welds is not None and self._last_opts == opts
                and self._last_plant == self._plant_key()):
            if snapshot:
                self._write_snapshot("焊接量", self._last_welds)
            else:
                self._on_welds_done(None, self._last_welds)
            return

        self._busy_on()
        token = self._tokens["stats"] = self._tokens["stats"] + 1
        info, plant = self.conn_info, self.plant
        q = self._q
        scope = self._scope_text() if snapshot else None
        # 焊接量统计口径固化进快照（分组模式/忽略壁厚/忽略直管焊缝）
        weld_opts = {"met": opts.met, "thk": opts.thk,
                     "skip_pipes": opts.skip_pipes}

        def work():
            try:
                db_r = db.open_db(info.get("provider", "MSSQL"), info["server"],
                                  plant["reports"], info["username"], info["password"])
                # 可选 MDB（model）连接：供 weld_core 解析读不到的代理端口。
                # 连接/打开失败不阻断主流程，weld_core 内部会自动降级。
                db_m = None
                if plant.get("model"):
                    try:
                        db_m = db.open_db(info.get("provider", "MSSQL"),
                                          info["server"], plant["model"],
                                          info["username"], info["password"])
                    except Exception:
                        db_m = None
                try:
                    welds = weld_core.extract_welds_via_sql(
                        db_r, plant["schema"], plant["catalog"], opts,
                        model_db=db_m, mdb_schema=plant.get("model") or None)
                    pipes = weld_core.aggregate_pipes(
                        db_r, plant["schema"], opts,
                        mdb_schema=plant.get("model") or None)
                    diag = ""
                    if not welds and opts.sys_oids:
                        # 空结果诊断：逐范围 oid 查 JNamedItem 匹配数，定位范围为何失效
                        parts = [f"当前 RDB: {plant['reports']}"]
                        for i, oid in enumerate(opts.sys_oids):
                            name = (opts.sys_names[i]
                                    if i < len(opts.sys_names) else "")
                            rows = db_r.query(
                                f"SELECT COUNT(*) FROM JNamedItem WHERE oid = "
                                f"'{weld_core.safe_sql(oid)}'")
                            parts.append(f"{name or oid}({oid}) -> "
                                         f"{rows[0][0] if rows else -1} 条匹配")
                        diag = ("统计结果为空。范围节点在 RDB(JNamedItem) 中的匹配情况：\n"
                                + "\n".join(parts))
                    for k, v in pipes.items():
                        welds[k] = welds.get(k, 0) + v
                    if not welds:
                        q.put(("welds_empty", token,
                               diag or "统计结果为空，且未设置系统范围。"))
                    elif snapshot:
                        q.put(("snap_done", token, snapshots.save(
                            "焊接量", plant["code"], scope, welds,
                            weld_opts=weld_opts)))
                    else:
                        q.put(("welds_done", token, welds))
                finally:
                    if db_m is not None:
                        db_m.close()
                    db_r.close()
            except Exception as ex:
                q.put(("welds_failed", token, str(ex)))
        self._start(work)

    def _on_welds_done(self, token, welds):
        self._busy_off()
        if token is not None and token != self._tokens["stats"]:
            return  # 旧线程的晚到结果：仅回收 busy 计数，不覆盖界面
        if not welds:
            messagebox.showwarning("提示", "统计结果为空。请检查系统范围与规格配置。",
                                   parent=self)
            return
        self._last_welds = welds
        self._last_opts = self._pending_opts
        self._last_plant = self._plant_key()
        default_name = f"{self.plant['code']}_焊接量.xlsx"
        start_dir = self.cfg.get("last_dir") or str(Path.home())
        path = filedialog.asksaveasfilename(
            parent=self, title="保存 Excel", initialdir=start_dir,
            initialfile=default_name, defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")])
        if not path:
            messagebox.showinfo(
                "提示",
                "已取消保存。统计结果已保留，再次点击「统计并生成 Excel」可直接导出"
                "（无需重新统计）。", parent=self)
            return
        try:
            import excel_out
            excel_out.write_weld_excel(welds, self.cmb_met.current(),
                                       self.chk_thk.get(), path)
            self.cfg["last_dir"] = str(Path(path).parent)
            db.save_config(self.cfg)
            self._last_welds = None  # 导出成功，清除缓存
            messagebox.showinfo(
                "成功", f"文件保存成功！\n{path}\n\n共 {len(welds)} 类焊口/管子行。",
                parent=self)
        except Exception as ex:
            messagebox.showerror(
                "导出失败",
                f"{ex}\n\n统计结果已保留，可再次点击「统计并生成 Excel」重新导出"
                "（无需重新统计）。", parent=self)

    def _on_welds_failed(self, token, msg):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        messagebox.showerror("统计失败", msg, parent=self)

    def _on_welds_empty(self, token, msg):
        self._busy_off()
        if token != self._tokens["stats"]:
            return
        messagebox.showwarning("提示", msg, parent=self)

    # ---------------- busy 计数 ----------------
    def _busy_buttons(self):
        return (self.btn_run, self.btn_refresh, self.btn_sys,
                self.btn_import, self.btn_export, self.btn_reread)

    def _busy_on(self):
        self._busy += 1
        if self._busy == 1:
            self.progress.start(12)
            for b in self._busy_buttons():
                b.configure(state="disabled")
            # 项目/材料范围切换会改变统计键：任务进行中必须禁用，否则
            # 晚到结果会用新状态命名导出旧行、并把旧行缓存成新键
            self.cmb_plant.configure(state="disabled")
            self.cmb_hidden.configure(state="disabled")
            self.cmb_report.configure(state="disabled")
            self.cmb_implied.configure(state="disabled")

    def _busy_off(self):
        self._busy = max(0, self._busy - 1)
        if self._busy == 0:
            self.progress.stop()
            # 项目未选时保持统计按钮禁用
            for b in self._busy_buttons():
                if b is not self.btn_run or self.plant:
                    b.configure(state="normal")
            # 下拉恢复 readonly（不能恢复 normal，否则可手输）
            self.cmb_plant.configure(state="readonly")
            self.cmb_report.configure(state="readonly")
            self.cmb_implied.configure(state="readonly")
            # 材料范围/分组选项仅带选项的报表可用（报表切走时保持禁用）
            opt_like = self._report() in ("管道材料", "螺栓", "垫片", "管夹")
            self.cmb_hidden.configure(
                state="readonly" if opt_like else "disabled")
            self.chk_by_line_cb.configure(
                state="normal" if opt_like else "disabled")


def main():
    app = App()
    app.mainloop()


if __name__ == "__main__":
    main()
